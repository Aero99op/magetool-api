"""
File Service - Magetool
Handles all file processing operations (PDFs, documents, etc.)
"""

from fastapi import UploadFile
from pathlib import Path
import uuid
import os

# Optional imports with fallbacks
try:
    from pypdf import PdfReader, PdfWriter
    PYPDF_AVAILABLE = True
except ImportError:
    PYPDF_AVAILABLE = False

try:
    from pdf2image import convert_from_path
    PDF2IMAGE_AVAILABLE = True
except ImportError:
    PDF2IMAGE_AVAILABLE = False

try:
    import pytesseract
    from PIL import Image
    TESSERACT_AVAILABLE = True
except ImportError:
    TESSERACT_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

TEMP_DIR = Path("./temp")

class FileService:
    """Service for file processing operations"""
    
    async def _save_uploaded_file(self, file: UploadFile) -> Path:
        """Save uploaded file to temp directory"""
        content = await file.read()
        suffix = Path(file.filename).suffix if file.filename else ''
        temp_path = TEMP_DIR / f"{uuid.uuid4()}{suffix}"
        temp_path.write_bytes(content)
        return temp_path
    
    async def convert_format(self, file: UploadFile, target_format: str) -> dict:
        """Convert file to target format"""
        target_format = target_format.lower().strip('.')
        temp_path = await self._save_uploaded_file(file)
        source_format = temp_path.suffix.lower().strip('.')
        
        try:
            # PDF conversions
            if source_format == 'pdf' and target_format in ['png', 'jpg', 'jpeg']:
                if not PDF2IMAGE_AVAILABLE:
                    return {"error": "pdf2image not installed"}
                
                images = convert_from_path(temp_path)
                output_files = []
                
                for i, img in enumerate(images):
                    output_filename = f"{uuid.uuid4()}_page{i+1}.{target_format}"
                    output_path = TEMP_DIR / output_filename
                    img.save(output_path, format=target_format.upper() if target_format != 'jpg' else 'JPEG')
                    output_files.append({
                        "filename": output_filename,
                        "page": i + 1,
                        "size": output_path.stat().st_size
                    })
                
                return {
                    "files": output_files,
                    "original": file.filename,
                    "page_count": len(images)
                }
            
            # Image to PDF
            elif source_format in ['png', 'jpg', 'jpeg', 'webp'] and target_format == 'pdf':
                from PIL import Image as PILImage
                img = PILImage.open(temp_path)
                if img.mode == 'RGBA':
                    img = img.convert('RGB')
                
                output_filename = f"{uuid.uuid4()}.pdf"
                output_path = TEMP_DIR / output_filename
                img.save(output_path)
                
                return {
                    "filename": output_filename,
                    "original": file.filename,
                    "size": output_path.stat().st_size
                }
            
            # DOCX to TXT
            elif source_format == 'docx' and target_format == 'txt':
                if not DOCX_AVAILABLE:
                    return {"error": "python-docx not installed"}
                
                doc = Document(temp_path)
                text = '\n'.join([para.text for para in doc.paragraphs])
                
                output_filename = f"{uuid.uuid4()}.txt"
                output_path = TEMP_DIR / output_filename
                output_path.write_text(text, encoding='utf-8')
                
                return {
                    "filename": output_filename,
                    "original": file.filename,
                    "size": output_path.stat().st_size
                }
            
            # TXT to DOCX
            elif source_format == 'txt' and target_format == 'docx':
                if not DOCX_AVAILABLE:
                    return {"error": "python-docx not installed"}
                
                text = temp_path.read_text(encoding='utf-8')
                doc = Document()
                for para in text.split('\n'):
                    doc.add_paragraph(para)
                
                output_filename = f"{uuid.uuid4()}.docx"
                output_path = TEMP_DIR / output_filename
                doc.save(output_path)
                
                return {
                    "filename": output_filename,
                    "original": file.filename,
                    "size": output_path.stat().st_size
                }
            
            else:
                return {"error": f"Conversion from {source_format} to {target_format} not supported"}
                
        finally:
            if temp_path.exists():
                os.remove(temp_path)
    
    async def ocr_scan(self, file: UploadFile) -> dict:
        """Extract text from file using OCR"""
        temp_path = await self._save_uploaded_file(file)
        source_format = temp_path.suffix.lower().strip('.')
        
        try:
            text = ""
            
            if source_format == 'pdf':
                if not PDF2IMAGE_AVAILABLE or not TESSERACT_AVAILABLE:
                    return {"error": "pdf2image and pytesseract required for PDF OCR"}
                
                images = convert_from_path(temp_path)
                for img in images:
                    text += pytesseract.image_to_string(img) + "\n"
            
            elif source_format in ['png', 'jpg', 'jpeg', 'webp', 'bmp', 'tiff']:
                if not TESSERACT_AVAILABLE:
                    return {"error": "pytesseract not installed"}
                
                img = Image.open(temp_path)
                text = pytesseract.image_to_string(img)
            
            else:
                return {"error": f"OCR not supported for {source_format} files"}
            
            return {
                "original": file.filename,
                "text": text.strip(),
                "char_count": len(text.strip())
            }
        finally:
            if temp_path.exists():
                os.remove(temp_path)
    
    async def edit_content(self, file: UploadFile, content: str) -> dict:
        """Edit file content (for text-based files)"""
        temp_path = await self._save_uploaded_file(file)
        source_format = temp_path.suffix.lower().strip('.')
        
        try:
            if source_format == 'txt':
                output_filename = f"{uuid.uuid4()}.txt"
                output_path = TEMP_DIR / output_filename
                output_path.write_text(content, encoding='utf-8')
                
                return {
                    "filename": output_filename,
                    "original": file.filename,
                    "size": output_path.stat().st_size
                }
            
            elif source_format == 'docx':
                if not DOCX_AVAILABLE:
                    return {"error": "python-docx not installed"}
                
                doc = Document()
                for para in content.split('\n'):
                    doc.add_paragraph(para)
                
                output_filename = f"{uuid.uuid4()}.docx"
                output_path = TEMP_DIR / output_filename
                doc.save(output_path)
                
                return {
                    "filename": output_filename,
                    "original": file.filename,
                    "size": output_path.stat().st_size
                }
            
            else:
                return {"error": f"Editing not supported for {source_format} files"}
                
        finally:
            if temp_path.exists():
                os.remove(temp_path)
    
    async def merge_pdfs(self, files: list[UploadFile]) -> dict:
        """Merge multiple PDFs into one"""
        if not PYPDF_AVAILABLE:
            return {"error": "pypdf not installed. Install with: pip install pypdf"}
        
        temp_paths = []
        
        try:
            merger = PdfWriter()
            
            for file in files:
                temp_path = await self._save_uploaded_file(file)
                temp_paths.append(temp_path)
                
                reader = PdfReader(temp_path)
                for page in reader.pages:
                    merger.add_page(page)
            
            output_filename = f"{uuid.uuid4()}_merged.pdf"
            output_path = TEMP_DIR / output_filename
            
            with open(output_path, 'wb') as f:
                merger.write(f)
            
            return {
                "filename": output_filename,
                "page_count": len(merger.pages),
                "files_merged": len(files),
                "size": output_path.stat().st_size
            }
        finally:
            for path in temp_paths:
                if path.exists():
                    os.remove(path)
    
    async def split_pdf(self, file: UploadFile, pages: str) -> dict:
        """Split PDF by page ranges (e.g., '1-3,5,7-9')"""
        if not PYPDF_AVAILABLE:
            return {"error": "pypdf not installed. Install with: pip install pypdf"}
        
        temp_path = await self._save_uploaded_file(file)
        
        try:
            reader = PdfReader(temp_path)
            total_pages = len(reader.pages)
            
            # Parse page ranges
            page_sets = []
            for part in pages.split(','):
                part = part.strip()
                if '-' in part:
                    start, end = map(int, part.split('-'))
                    page_sets.append(list(range(start, min(end + 1, total_pages + 1))))
                else:
                    page_num = int(part)
                    if 1 <= page_num <= total_pages:
                        page_sets.append([page_num])
            
            output_files = []
            
            for i, page_nums in enumerate(page_sets):
                writer = PdfWriter()
                for page_num in page_nums:
                    writer.add_page(reader.pages[page_num - 1])
                
                output_filename = f"{uuid.uuid4()}_split_{i+1}.pdf"
                output_path = TEMP_DIR / output_filename
                
                with open(output_path, 'wb') as f:
                    writer.write(f)
                
                output_files.append({
                    "filename": output_filename,
                    "pages": page_nums,
                    "size": output_path.stat().st_size
                })
            
            return {
                "files": output_files,
                "original": file.filename,
                "original_pages": total_pages
            }
        finally:
            if temp_path.exists():
                os.remove(temp_path)
